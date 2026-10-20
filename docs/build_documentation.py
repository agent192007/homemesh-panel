#!/usr/bin/env python3
"""Build docs/HomeMesh_Panel_Project_Documentation.docx.

The handover document a teammate builds the presentation from. Written
for someone who has not worked on the code. Rebuild it from this script;
do not hand-edit the .docx, because the next build overwrites it.

    python docs/build_documentation.py            # build, then Word fills the contents page
    python docs/build_documentation.py --no-word  # build only (contents page fills on open)

Needs python-docx. The Word step (Windows, Microsoft Word installed)
updates the table of contents, saves, and exports a PDF beside the .docx.

Figures come from scripts, not from hand edits:
  docs/figures/generate_diagrams.py   arch_overall / arch_hardware / arch_software
  docs/figures/paper_block_diagram.py docs/figures/paper_figures.py   img/paper_*
  designer/tests/render_png.py, theme_sheet.py   panel renders (the firmware's own UI)
The device table is read from include/ha_layout.json at build time, so it
always matches the layout in the repository.
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FIG = HERE / "figures"
IMG = FIG / "img"
OUT = HERE / "HomeMesh_Panel_Project_Documentation.docx"

UPDATED = "20 October 2026"
SUBMISSION = "21 October 2026"
TEAM = [("Garima Nidhi", "24BCE1782"), ("Taniska Routray", "24BCE1578"),
        ("Bharadhwaj B", "24BCE1959"), ("Pranesh S", "24BCE5138")]
GUIDE = "Suhasini S"

FONT = "Calibri"
MONO = "Consolas"
INK = RGBColor(0x1F, 0x23, 0x2B)
HEAD = RGBColor(0x1A, 0x2B, 0x4C)
GREY = RGBColor(0x5F, 0x63, 0x6B)
HEAD_FILL = "E9E9EF"
NOTE_FILL = "FFF6E5"
INFO_FILL = "EAF1FB"
CODE_FILL = "F4F4F6"
BORDER = "B8BCC6"

# ----------------------------------------------------------------------
#  Low-level helpers
# ----------------------------------------------------------------------
doc = Document()


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def borders(table, color=BORDER, size=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    trPr.append(e)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    trPr.append(e)


def fixed_widths(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for row in table.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)


TOKEN = re.compile(r"(<b>.*?</b>|<i>.*?</i>|<c>.*?</c>)")


def add_runs(p, text, size=None, color=None, bold=False, italic=False):
    """Inline markup: <b>bold</b>, <i>italic</i>, <c>code</c>."""
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, mono = bold, italic, False
        if part.startswith("<b>"):
            part, b = part[3:-4], True
        elif part.startswith("<i>"):
            part, i = part[3:-4], True
        elif part.startswith("<c>"):
            part, mono = part[3:-4], True
        r = p.add_run(part)
        r.bold, r.italic = b, i
        if mono:
            r.font.name = MONO
            r.font.size = Pt((size or 10.5) - 1)
        elif size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return p


def para(text, size=None, color=None, align=None, after=6, before=0, bold=False, italic=False, keep=False):
    p = doc.add_paragraph()
    add_runs(p, text, size, color, bold, italic)
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    if align:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    return p


def h1(text, new_page=False):
    p = doc.add_heading(text, level=1)
    if new_page:
        p.paragraph_format.page_break_before = True
    return p


def h2(text):
    return doc.add_heading(text, level=2)


def h3(text):
    return doc.add_heading(text, level=3)


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        add_runs(p, it)
        p.paragraph_format.space_after = Pt(3)


def numbered(items):
    # Numbered by hand: Word's List Number continues across lists, so every
    # list after the first would start where the previous one stopped.
    for n, it in enumerate(items, 1):
        p = doc.add_paragraph()
        add_runs(p, f"{n}.\t" + it)
        pf = p.paragraph_format
        pf.left_indent = Cm(0.75)
        pf.first_line_indent = Cm(-0.75)
        pf.tab_stops.add_tab_stop(Cm(0.75))
        pf.space_after = Pt(3)


def table(headers, rows, widths, size=9.5, header=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    cell_margins(t)
    if header and headers:
        r = t.add_row()
        repeat_header(r)
        for c, h in zip(r.cells, headers):
            shade(c, HEAD_FILL)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            add_runs(c.paragraphs[0], h, size=size, bold=True)
    for row in rows:
        r = t.add_row()
        no_split(r)
        for c, v in zip(r.cells, row):
            lines = v.split("\n") if isinstance(v, str) else v
            for k, ln in enumerate(lines):
                pp = c.paragraphs[0] if k == 0 else c.add_paragraph()
                pp.paragraph_format.space_after = Pt(0)
                add_runs(pp, ln, size=size)
    fixed_widths(t, widths)
    spacer(4)
    return t


def spacer(pt=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run()
    r.font.size = Pt(pt)
    return p


def callout(title, lines, fill=NOTE_FILL):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, color="D9C9A3" if fill == NOTE_FILL else "A9BFDF", size=6)
    cell_margins(t, 120, 120, 160, 160)
    no_split(t.rows[0])
    c = t.rows[0].cells[0]
    shade(c, fill)
    p = c.paragraphs[0]
    add_runs(p, title, size=10.5, bold=True)
    p.paragraph_format.space_after = Pt(3)
    for ln in ([lines] if isinstance(lines, str) else lines):
        q = c.add_paragraph()
        add_runs(q, ln, size=10)
        q.paragraph_format.space_after = Pt(2)
    fixed_widths(t, [16.0])
    spacer(6)


def code(text, size=8.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, color="D5D7DD")
    cell_margins(t, 100, 100, 160, 160)
    no_split(t.rows[0])
    c = t.rows[0].cells[0]
    shade(c, CODE_FILL)
    lines = text.split("\n")
    for k, ln in enumerate(lines):
        p = c.paragraphs[0] if k == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln if ln else " ")
        r.font.name = MONO
        r.font.size = Pt(size)
    fixed_widths(t, [16.0])
    spacer(6)


FIGNO = [0]


def figure(path, width_cm, caption):
    FIGNO[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    para(f"<b>Figure {FIGNO[0]}.</b> {caption}", size=9, color=GREY,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    return FIGNO[0]


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc():
    p = doc.add_paragraph()
    r = p.add_run()
    for kind, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
    r2 = p.add_run("Right-click here and choose Update Field to fill in the contents.")
    r2.italic = True
    r3 = p.add_run()
    f = OxmlElement("w:fldChar")
    f.set(qn("w:fldCharType"), "end")
    r3._r.append(f)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HomeMesh Panel — Project Documentation   |   page ")
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    r = p.add_run()
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)


# ----------------------------------------------------------------------
#  Document setup
# ----------------------------------------------------------------------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))

st = doc.styles["Normal"]
st.font.name = FONT
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.font.size = Pt(10.5)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12

for name, size, before, after in (("Heading 1", 16, 18, 8), ("Heading 2", 12.5, 12, 4), ("Heading 3", 11, 8, 3)):
    s = doc.styles[name]
    s.font.name = FONT
    rf = s.element.rPr.rFonts
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = HEAD
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number"):
    doc.styles[name].font.name = FONT
    doc.styles[name].font.size = Pt(10.5)

# ----------------------------------------------------------------------
#  Data read from the repository
# ----------------------------------------------------------------------
layout = json.loads((ROOT / "include" / "ha_layout.json").read_text(encoding="utf-8"))
areas = layout["areas"]
ents = layout["entities"]
sensors = layout.get("sensors", [])
actions = layout.get("actions", [])
KIND_WORDS = {"Toggle": "on/off", "Dimmable": "brightness", "RGB": "brightness and colour"}


def ent_desc(e):
    dom = e["id"].split(".")[0]
    what = {"light": "light", "switch": "switch", "fan": "fan", "input_boolean": "helper switch"}.get(dom, dom)
    return f"{e['name']} ({what}, {KIND_WORDS.get(e['kind'], e['kind'])})"


# ======================================================================
#  TITLE PAGE
# ======================================================================
spacer(60)
para("HOMEMESH PANEL", size=28, bold=True, color=HEAD, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("A Wall-Mounted Smart Home Control Panel using ESP32 and the Home Assistant WebSocket API",
     size=13, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
para("CSE3006 — Embedded System Design", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Final Review — Project Documentation", size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=28)

table(["", ""], [
    ["Project title", "HomeMesh Panel — local-first smart home control panel"],
    ["Institution", "Vellore Institute of Technology, Chennai"],
    ["Course", "CSE3006 — Embedded System Design"],
    ["Team members", "\n".join(n for n, _ in TEAM)],
    ["Register numbers", "\n".join(f"{r}  ({n})" for n, r in TEAM)],
    ["Faculty guide", GUIDE],
    ["Date of submission", SUBMISSION],
    ["Repository", "github.com/agent192007/homemesh-panel (private)"],
], [4.5, 10.5], size=10, header=False)

spacer(30)
para(f"This version: {UPDATED}. It replaces the version of 17 September, which described the project "
     "before the layout moved onto the network and before the desktop designer existed.",
     size=9, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

# ======================================================================
#  CONTENTS + HOW TO USE
# ======================================================================
# Not a Heading 1, so the contents page does not list itself.
p = para("Table of Contents", size=16, bold=True, color=HEAD, after=8)
p.paragraph_format.page_break_before = True
toc()

h1("How to Use This Document", new_page=True)
para("This document is written for a team member who has not worked on the code, and who is building the "
     "presentation from it. It assumes no prior knowledge of the project and explains each term the first "
     "time it appears.")
para("Sections 1 to 12 follow the structure the course expects and map almost one-to-one onto slides. "
     "Section 13 has the measured results and the problems solved along the way. <b>Appendix A</b> is a "
     "slide-by-slide plan with the figure to use on each slide and likely questions with their answers. "
     "<b>Appendix B</b> is an internal note for the team only: what can be claimed, what cannot, and the "
     "checklist for the morning of the review.")
para("Every figure in this document is also a separate image file, listed in Appendix A, and can be "
     "dropped straight onto a slide. Every number was measured on the running system, not estimated.")
callout("Read Appendix B before presenting",
        ["Appendix B lists what is not finished and the specific claims that must not be made during the "
         "review. Everything else in this document is safe to present."])
callout("What changed since the 17 September version", [
    "• The device list is <b>no longer compiled into the firmware</b>. It is designed on a laptop and sent to "
    "the panel over the network, where it is stored. No cable, no reflashing. Proven on the board on 17 October.",
    "• A new <b>desktop designer app</b> replaces the command-line tool. Its preview is the panel's own "
    "interface code running on the PC, not a drawing of it.",
    "• The panel now has <b>device icons</b>, a <b>new tile design</b>, <b>custom device names</b>, and "
    "<b>six colour themes</b>, all chosen in the designer.",
    "• The panel takes device state over a <b>single subscription</b> (<c>subscribe_entities</c>) instead of "
    "one request per device. This fixed a connection-loss bug, described in Section 13.4.",
    "• A <b>Wi-Fi fault</b> that looked like random firmware failure was traced to the laptop's Wi-Fi card. "
    "It makes a good story for the review, also in Section 13.4.",
], fill=INFO_FILL)

# ======================================================================
#  1. ABSTRACT
# ======================================================================
h1("1. Abstract", new_page=True)
para("HomeMesh Panel is a wall-mounted touchscreen that controls the devices in a home automation system. "
     "It is built on an ESP32 microcontroller with a 3.2-inch capacitive touchscreen, and it communicates with "
     "a Home Assistant server over the local wireless network using the Home Assistant WebSocket API.")
para("The panel shows lights, switches, fans and environmental sensors as touchable tiles, on pages that "
     "correspond to rooms, plus a Home page of favourites. A tap switches a device; a detail panel sets "
     "brightness and colour. State travels in both directions: a change made anywhere else in the system "
     "appears on the panel within milliseconds, because the server pushes every change to it.")
para("The central engineering problem is memory. Asking Home Assistant for the state of every device returned "
     "<b>57,671 bytes for the 123 devices</b> of the test installation, more than the panel's entire 40 KB "
     "graphics memory, on a chip with no external RAM. The project solves this by splitting the work across two "
     "machines: a desktop application performs discovery and designs the panel's layout, and the panel "
     "subscribes only to the devices in that layout. Its initial state for seven devices arrives in a single "
     "<b>2,403-byte</b> message.")
para("The layout — which devices, on which pages, with which names, icons and colours — is data, not code. "
     "It is sent to the panel through Home Assistant and stored in the panel's non-volatile memory, so the panel "
     "is reconfigured in seconds without reprogramming. The desktop designer previews the result using the "
     "panel's own interface code compiled to run on a PC.")
para("The system runs entirely on the local network. It needs no cloud service, no vendor app and no internet "
     "connection.")

# ======================================================================
#  2. INTRODUCTION
# ======================================================================
h1("2. Introduction")
para("Home automation systems have become common, but the way people interact with them has not kept pace. "
     "Most are controlled through a smartphone app, so switching on a light means unlocking a phone, finding "
     "an app, waiting for it to connect and navigating to the right screen. That is slower and less reliable "
     "than the wall switch it replaced, and it leaves out visitors and anyone without the phone.")
para("A wall-mounted control panel removes that friction. It is always on, always connected and always in "
     "the same place. It behaves like a light switch, but it can control any device in the system and show "
     "information a switch cannot, such as room temperature and humidity.")
para("Commercial wall panels exist, but they are expensive, usually tied to one manufacturer, and often "
     "depend on a cloud service. A cloud dependency means a device two metres away stops responding when the "
     "internet connection fails, and that a record of household activity leaves the home.")
para("This project builds such a panel from inexpensive, widely available parts, controls an open-source "
     "home automation platform, and works entirely within the local network.")
h2("2.1 Home Assistant")
para("Home Assistant is an open-source home automation platform: the software that actually talks to smart "
     "devices. It supports a very large number of device types through components called <i>integrations</i> "
     "— one for Zigbee devices, another for Z-Wave, another for Wi-Fi devices, and so on.")
para("The important consequence is that Home Assistant presents all of these devices through one uniform "
     "interface, whatever the underlying wireless technology. Each controllable thing is called an "
     "<i>entity</i>, with an identifier such as <c>light.bed_light</c>. The panel is a client of Home "
     "Assistant: it never talks to bulbs or switches directly, it asks Home Assistant to.")
h2("2.2 Literature survey")
para("The project draws on four strands of published work. The full references are in Section 16.1.")
table(["Area", "What the literature says", "What HomeMesh takes from it"], [
    ["Edge and local processing [1], [2]",
     "Processing inside the home, rather than in a vendor cloud, reduces latency, keeps working without "
     "the internet and keeps household data local.",
     "No cloud anywhere in the path. Server, network and panel are all on the premises."],
    ["Interoperability [3], [4]",
     "Device fragmentation across Zigbee, Z-Wave, Wi-Fi and vendor clouds has been the central problem of "
     "home automation for two decades; Matter is the latest attempt to solve it.",
     "Talk to Home Assistant, which already integrates devices from many ecosystems, instead of to devices."],
    ["ESP32 home automation [5], [6], [7]",
     "Typical ESP32 systems switch relays and read sensors wired to the board, and are controlled from a "
     "web page or cloud dashboard.",
     "The ESP32 is the controller, not the controlled: one panel for every device the server knows."],
    ["Accessible fixed panels [8]",
     "A fixed panel gives every occupant, including people without a smartphone, control of the room they "
     "are in.",
     "The motivation for a wall-mounted panel at all."],
    ["IoT protocols [9]–[12]",
     "MQTT is the dominant IoT messaging protocol; surveys compare it with CoAP, HTTP, AMQP and others.",
     "Why the first version used MQTT, and the reach problem that made it a dead end (Section 5.1)."],
    ["MQTT vs WebSocket [13]–[16]",
     "WebSocket gives the lowest latency for interactive, request-driven exchanges; MQTT is more economical "
     "for streaming telemetry.",
     "A control panel is interactive by nature, which favours WebSocket."],
    ["ESP32 power [17]", "Measured power consumption of the ESP32 in low-power applications.",
     "Not a goal here: the panel is mains powered, so Wi-Fi power saving is switched off for responsiveness."],
], [3.4, 6.6, 6.0], size=9)

# ======================================================================
#  3. PROBLEM STATEMENT
# ======================================================================
h1("3. Problem Statement")
para("Design and build a wall-mounted touchscreen control panel for a home automation system that satisfies "
     "these conditions:")
numbered([
    "It must control real devices of different types through one interface, without being restricted to one "
    "wireless technology or one manufacturer.",
    "It must show the true current state of each device, and update when that state is changed by any other "
    "means.",
    "It must work entirely on the local network, with no cloud service in the control path.",
    "It must run on a low-cost microcontroller with severely limited memory, which rules out the "
    "straightforward approach of downloading the entire system state.",
    "It must be reconfigurable — which devices, where, how they look — by someone with no programming "
    "toolchain.",
    "The interface must stay responsive to touch at all times, including while the network is busy.",
])

# ======================================================================
#  4. OBJECTIVES
# ======================================================================
h1("4. Objectives")
numbered([
    "To interface an ESP32 microcontroller with a 320 × 240 colour display and a capacitive touch panel, and "
    "to identify the controller variants actually fitted to the board.",
    "To implement a WebSocket client on the ESP32 that authenticates with Home Assistant and holds one "
    "persistent connection to it.",
    "To receive device state pushed by the server and reflect it on the display without polling.",
    "To send control commands from the panel so that devices can be switched, dimmed and coloured from the "
    "touchscreen.",
    "To design a memory strategy that lets a device with no external RAM take part in a system whose full "
    "state is far larger than its memory.",
    "To build a multi-page, room-based interface that stays responsive during network activity.",
    "To make the panel's layout data rather than code: authored on a PC, delivered over the network, stored "
    "on the panel, with no reflashing.",
    "To build a desktop designer that discovers devices from the server and previews the panel exactly as it "
    "will appear.",
])

# ======================================================================
#  5. EXISTING SYSTEMS
# ======================================================================
h1("5. Existing Systems and Their Limitations")
para("Three categories of existing solution were examined before the architecture was chosen.")
table(["Approach", "How it works", "Limitations"], [
    ["Manufacturer wall panels", "A sealed product sold as part of one manufacturer's ecosystem.",
     "High cost. Controls only that manufacturer's devices. Usually needs a cloud account, so it fails when "
     "the internet fails. Not extensible."],
    ["Tablet running a web dashboard", "A consumer tablet on the wall showing the home automation web "
     "interface in a browser.",
     "Expensive for what it does. Slow to boot. The battery degrades on permanent charge. A full operating "
     "system to maintain and update."],
    ["Microcontroller panel using MQTT", "An ESP32 connected to an MQTT message broker, publishing and "
     "subscribing to topics that represent devices.",
     "Only reaches devices that publish themselves to the broker. Large classes of device are invisible. "
     "See Section 5.1."],
], [3.6, 5.6, 6.8], size=9)
h2("5.1 Why the MQTT approach was rejected")
para("MQTT is a lightweight publish-subscribe messaging protocol widely used in embedded systems, and it was "
     "the approach used in the first version of this project. It was abandoned, and the reason is the most "
     "important design decision in the work.")
para("MQTT can only reach devices that publish themselves onto a message broker. A device connected to Home "
     "Assistant through a non-MQTT integration — a Philips Hue bulb, a Z-Wave switch, or any device added "
     "through the ZHA Zigbee integration — exists entirely inside Home Assistant and has no MQTT topic at all. "
     "A panel speaking MQTT cannot see such a device, and so cannot control it.")
para("This is not just a limitation; it is unpredictable for the user. <b>The same physical Zigbee bulb</b> is "
     "reachable over MQTT if it was added through Zigbee2MQTT, and completely invisible if it was added through "
     "ZHA. The hardware is identical. Whether the panel can control it depends on a setup decision the user had "
     "no reason to think about. A panel whose coverage depends on how each device happened to be added is not "
     "a usable product.")
h2("5.2 The three designs the project went through")
table(["Design", "Devices reached", "How state arrived", "Why it was replaced"], [
    ["1. MQTT broker (Sept)", "Only those publishing to the broker", "Retained MQTT messages",
     "Most Home Assistant devices invisible"],
    ["2. WebSocket + REST (Sept – 17 Oct)", "Every Home Assistant entity",
     "One HTTP request per device, then WebSocket events",
     "The requests blocked the connection and caused reconnect loops (Section 13.4)"],
    ["3. WebSocket only (18 Oct, final)", "Every Home Assistant entity",
     "One subscription: a snapshot, then only changes", "—"],
], [4.2, 3.6, 4.4, 3.8], size=9)

# ======================================================================
#  6. PROPOSED SYSTEM
# ======================================================================
h1("6. Proposed System")
para("The proposed system replaces MQTT with the Home Assistant WebSocket API — the same interface Home "
     "Assistant's own web interface uses. It has two properties MQTT cannot offer:")
bullets([
    "<b>Uniform reach.</b> Every entity Home Assistant knows is available, whichever integration created it. "
    "A Zigbee bulb, a Z-Wave switch, a Wi-Fi plug and a software-only helper are all addressed the same way.",
    "<b>One control command.</b> A single message type, <c>call_service</c>, operates all of them. Switching a "
    "Zigbee bulb and switching a Wi-Fi plug are the same message with a different target.",
])
para("A <i>WebSocket</i> is a network connection that stays open. With an ordinary web request the client "
     "asks, the server answers and the connection closes; a WebSocket stays established so either side can "
     "send at any time. That is what lets the server push a change to the panel the moment it happens, instead "
     "of the panel repeatedly asking (<i>polling</i>), which wastes power and network capacity and adds delay.")
h2("6.1 The memory constraint and the split architecture")
para("Adopting the WebSocket API brings a problem. The natural first step for any client is to ask for the "
     "state of everything, with a request called <c>get_states</c>. On the test installation that returned "
     "<b>57,671 bytes for 123 entities</b> — about 469 bytes each. A real house with a few hundred devices "
     "would be hundreds of kilobytes.")
para("The ESP32-WROOM-32 has no external RAM (PSRAM), and its graphics library works inside a fixed 40 KB "
     "pool. Even the small test installation's full state is bigger than that. This is not a problem that more "
     "careful code can solve; the data is simply larger than the memory.")
para("So the work is split across two machines. <b>This split is the central design idea of the project:</b>")
bullets([
    "<b>The PC does discovery.</b> The desktop designer connects to Home Assistant, calls <c>get_states</c>, "
    "shows every entity grouped by room, and the user picks the handful the panel should show.",
    "<b>The panel never does discovery.</b> It subscribes to exactly the entities in its layout and sends "
    "commands only to them. Every message it handles is small and bounded by construction.",
])
para("Measured effect: the panel's initial state for its seven entities arrives in one <b>2,403-byte</b> "
     "message — 24 times smaller — and that size grows with the layout, not with the size of the house.")
h2("6.2 The layout is data, not code")
para("In the first versions, the list of devices was compiled into the firmware: changing it meant rebuilding "
     "and reflashing over USB. It is now data. The designer produces a small JSON document (about 700 bytes) "
     "naming the rooms, the devices, their names, icons and kinds, the sensors, and optional theme colours. It "
     "is sent to the panel through Home Assistant and stored in the panel's non-volatile memory (NVS — flash "
     "memory that survives power loss). The firmware never changes; only the layout does.")
code('{"v":1,\n'
     ' "areas":["Bedroom","Kitchen","Living Room"],\n'
     ' "entities":[\n'
     '   {"id":"light.bed_light", "name":"Bed Light", "kind":"RGB",\n'
     '    "area":0, "fav":true, "icon":"light"},\n'
     '   ...],\n'
     ' "sensors":[{"id":"sensor.outside_temperature", "area":2, "hum":false}, ...],\n'
     ' "theme":{"bg":"#F2F3F7", ...}     <- optional: only colours that differ\n'
     '}')

# ======================================================================
#  7. SYSTEM ARCHITECTURE
# ======================================================================
h1("7. System Architecture", new_page=True)
para("The architecture is shown in three views, following the course convention: an overall system diagram "
     "of the major blocks and the data between them, a hardware diagram of the physical components and their "
     "electrical interfaces, and a software diagram of the program structure. Section 7.4 gives a combined "
     "block diagram used in the paper.")
h2("7.1 Overall system diagram")
figure(FIG / "arch_overall.png", 16, "Overall system architecture. The setup path runs on the laptop whenever "
       "the layout changes; the runtime path runs continuously on the panel.")
para("<b>Setup path (top).</b> The designer asks Home Assistant for every entity, the user chooses and arranges "
     "them, and the designer sends the resulting layout as a Home Assistant event called "
     "<c>homemesh_layout</c>. Home Assistant relays it to the panel over the connection the panel already "
     "holds. The panel validates it, stores it in NVS and restarts into it. No web server runs on the panel, "
     "and neither side ever needs the panel's IP address — both only need to know where Home Assistant is.")
para("<b>Runtime path (bottom).</b> Physical devices are controlled by Home Assistant through their "
     "integrations. The panel holds one WebSocket connection to Home Assistant over the laptop's Wi-Fi "
     "hotspot, receives state through it and sends commands back through it. The user interacts by touch.")
h2("7.2 Hardware diagram")
figure(FIG / "arch_hardware.png", 15.5, "Hardware architecture: the microcontroller and its peripheral "
       "interfaces.")
para("The ESP32-WROOM-32 module is the centre of the system. The display is driven over a four-wire SPI bus at "
     "27 MHz; the touch panel is read over I²C at 400 kHz; the backlight is switched by one GPIO pin. A CH340 "
     "chip provides the USB-to-serial link used for programming and for the diagnostic log. The 2.4 GHz Wi-Fi "
     "radio is built into the module.")
h2("7.3 Software diagram")
figure(FIG / "arch_software.png", 16, "Software architecture. The two processor cores are separated by a "
       "queue; the layout is read from NVS at boot.")
para("The software is split across the ESP32's two cores. Core 1 runs the user interface; core 0 runs the "
     "network client. They never call into each other; they communicate through a FreeRTOS queue — a "
     "thread-safe buffer one task writes and the other reads. This separation is why the interface stays "
     "smooth while the network is busy (Section 10.2). The bottom row shows what drives the firmware: the "
     "layout from NVS (with a compiled-in fallback), the theme colours, and the compiled configuration.")
h2("7.4 Block diagram")
figure(IMG / "paper_block_diagram.png", 16, "System block diagram, as used in the IEEE paper.")

# ======================================================================
#  8. HARDWARE
# ======================================================================
h1("8. Hardware Description", new_page=True)
para("The project uses a <b>Probots ESP32-2432S032</b> development board, from a family nicknamed the "
     "\"Cheap Yellow Display\", which puts an ESP32 module, a colour display and a touch panel on one board.")
h2("8.1 Components")
table(["Component", "Part", "Specification and role"], [
    ["Microcontroller", "ESP32-WROOM-32", "Dual-core 32-bit processor at 240 MHz. 520 KB on-chip SRAM, no "
     "external PSRAM. 4 MB flash. Integrated 2.4 GHz Wi-Fi (no 5 GHz). Runs the whole application."],
    ["Display", "ST7789V controller\n3.2\" 240 × 320 IPS", "Used in landscape at 320 × 240, 65,536 colours "
     "(16-bit RGB565). SPI at 27 MHz. Needs colour inversion on and BGR colour order."],
    ["Touch panel", "Goodix GT911", "Projected capacitive. I²C at 400 kHz, address 0x5D."],
    ["USB interface", "CH340", "USB-to-serial for firmware upload and the diagnostic log."],
    ["Storage", "4 MB SPI flash", "3 MB application partition, plus NVS where the pushed layout is stored."],
    ["microSD slot", "128 GB card fitted", "Physically present but <b>not used and not yet tested</b>. "
     "Candidate uses: logging, assets (Section 14)."],
    ["Power", "5 V over USB-C", "On-board regulator supplies 3.3 V."],
], [3.0, 3.8, 9.2], size=9)
h2("8.2 Pin assignment")
para("These values are specific to this board. They are set as build flags in <c>platformio.ini</c>, not in "
     "the display library's own configuration file.")
table(["Signal", "Pin", "Interface", "Function"], [
    ["MOSI", "13", "SPI", "Serial data to the display"],
    ["SCLK", "14", "SPI", "Serial clock for the display"],
    ["CS", "15", "SPI", "Display chip select"],
    ["DC", "2", "SPI", "Data / command select"],
    ["MISO", "12", "SPI", "Serial data from the display"],
    ["RST", "−1", "SPI", "Not connected; reset handled internally"],
    ["BL", "27", "GPIO", "Backlight enable, active high"],
    ["SDA", "33", "I²C", "Touch panel data"],
    ["SCL", "32", "I²C", "Touch panel clock"],
], [2.4, 1.6, 2.4, 9.6], size=9)
h2("8.3 Batch variation in this board family")
para("Boards in this family are sold under the same name while carrying different display and touch "
     "controllers from batch to batch, so published tutorials are often wrong for any particular unit. The "
     "parts on this board were identified experimentally: an I²C scanning program found the touch controller "
     "and its address, and a display test program confirmed the display controller by seeing which "
     "initialisation produced a correct image.")
para("The result: an <b>ST7789V</b> display controller, not the ILI9341 most tutorials assume, and a <b>GT911</b> "
     "capacitive touch controller. These symptoms identify a colour fault immediately if one ever recurs:")
table(["Symptom", "Cause"], [
    ["Red and blue swapped; green, white and black correct", "Colour order wrong — the panel is BGR, not RGB"],
    ["All colours inverted (red shows as cyan)", "Display inversion setting wrong for this panel"],
    ["Image shifted with a noise band along one edge", "Memory offsets wrong. This board needs none, so it "
     "did not occur"],
], [8.0, 8.0], size=9)
para("Because of this variation, all hardware-specific code lives in two driver files. The interface and "
     "network code never see it — a decision that later made the transport rewrite cheap.")

# ======================================================================
#  9. SOFTWARE
# ======================================================================
h1("9. Software Description", new_page=True)
h2("9.1 Panel firmware: environment and libraries")
table(["Layer", "Choice", "Version", "Reason"], [
    ["Build system", "PlatformIO", "—", "Libraries, build configurations and upload. No definition exists "
     "for this board, so the generic esp32dev target is used."],
    ["Framework", "Arduino for ESP32, on FreeRTOS", "—", "Widest library support; almost all reference "
     "material for this board assumes it."],
    ["Graphics", "LVGL", "8.3.11", "Buttons, sliders, colour wheel, pages. Pinned to 8.3 deliberately: "
     "version 9 changed the driver interfaces and almost every example for this board is 8.3."],
    ["Display driver", "TFT_eSPI", "2.5.43", "Fast SPI display driver with ST7789V support."],
    ["WebSocket", "links2004 WebSockets", "2.7.3", "Maintained WebSocket client for the ESP32."],
    ["JSON", "ArduinoJson", "6.21.6", "Filtered, fixed-size parsing with bounded memory."],
    ["Icons", "Material Design Icons as LVGL fonts", "16 glyphs", "Drawn like text from flash: about 7.6 KB "
     "of flash and nothing from the graphics memory."],
    ["Partitions", "huge_app.csv", "—", "3 MB application partition instead of 1.3 MB."],
], [2.6, 3.8, 1.8, 7.8], size=9)
h2("9.2 Desktop designer: environment and libraries")
table(["Layer", "Choice", "Reason"], [
    ["Application shell", "Electron", "A desktop app that runs on any Windows PC as a single portable .exe; "
     "nothing to install."],
    ["Interface", "React 19 + TypeScript", "Typed, component-based interface; the layout format is checked at "
     "compile time."],
    ["Build", "Vite", "Fast builds and hot reload during development."],
    ["Panel simulator", "Emscripten (C/C++ → WebAssembly)", "Compiles the panel's own interface code, LVGL "
     "and fonts to run inside the app. See Section 11.2."],
    ["Home Assistant client", "Browser WebSocket", "The same protocol the panel uses, without its memory "
     "limits."],
    ["Reference tools", "Python 3: ha_discover.py, preflight.ps1", "The layout format is defined once in "
     "Python; the app's version is tested against it byte for byte."],
], [3.4, 4.6, 8.0], size=9)
h2("9.3 Source code organisation")
para("Hardware details are confined to two files and never appear in the interface or network code.")
table(["File", "Responsibility"], [
    ["src/main.cpp", "Entry point. Starts the drivers and the network task, runs the graphics loop, and reports "
     "memory use and graphics failures."],
    ["src/display_driver.*", "Display start-up and the callback that sends rendered pixels to the panel."],
    ["src/touch_driver.*", "Touch controller and the callback that reports touches, including rotation."],
    ["src/ha_client.*", "The Home Assistant WebSocket client: connect, authenticate, subscribe, send commands, "
     "receive layouts."],
    ["src/layout_store.*", "The layout: read from NVS at boot, validate, store new ones, fall back to the "
     "compiled layout."],
    ["src/panel_theme.*", "The nine theme colours, defaulted and overridden by the layout."],
    ["src/ui/dashboard.*", "All screens: pages, tiles, the detail panel, the automations screen."],
    ["src/ui/panel_icons.*, fonts/", "Icon name → glyph, and the icon fonts."],
    ["src/panel_model.h", "The contract between transport and interface (Section 9.4)."],
    ["src/mqtt_client.*, fake_client.cpp", "The original MQTT transport and an offline simulated one. Unused "
     "in the demo, kept as fallbacks."],
    ["include/ha_layout.h / .json", "The compiled-in fallback layout, used when nothing has been pushed."],
    ["include/secrets.h", "Wi-Fi name and password, server address, access token. Never committed."],
    ["designer/", "The desktop designer app (Section 11)."],
    ["sim/", "Builds the panel's interface to WebAssembly for the designer."],
    ["tools/", "ha_discover.py (discovery and layout format), preflight.ps1 (demo-day checks), the older "
     "tkinter designer."],
    ["docs/", "The IEEE paper, this document and the scripts that draw every figure."],
], [5.2, 10.8], size=9)
h2("9.4 The transport interface")
para("<c>panel_model.h</c> is the boundary between the network code and the interface code. The network side "
     "knows nothing about buttons; the interface knows nothing about sockets. They share only three "
     "definitions:")
table(["Definition", "Meaning"], [
    ["EntityKind", "What controls a device gets. Toggle: on/off. Dimmable: adds brightness. RGB: adds colour."],
    ["LinkState", "Connection condition — no Wi-Fi, no server, or online. Drives the status dot."],
    ["PanelEvent", "A small fixed-size record of one change: which device, what changed, the new value. "
     "Copied through the queue."],
], [3.4, 12.6], size=9)
para("The value of this boundary showed when MQTT was replaced by the WebSocket client: <b>this file did not "
     "change at all</b>. The new transport offered the same three functions as the old one — start, next "
     "event, send — so the main program and the interface barely changed. A potential rewrite was contained "
     "to one file.")

# ======================================================================
#  10. IMPLEMENTATION
# ======================================================================
h1("10. Implementation", new_page=True)
h2("10.1 The communication protocol")
para("From the panel's point of view, the whole of a home automation system reduces to five exchanges on one "
     "connection:")
code("1.  Connect     ws://192.168.137.1:8123/api/websocket\n"
     "                server:  {\"type\":\"auth_required\"}\n"
     "\n"
     "2.  Log in      {\"type\":\"auth\",\"access_token\":\"<long-lived token>\"}\n"
     "                server:  {\"type\":\"auth_ok\"}\n"
     "\n"
     "3.  Subscribe   {\"id\":1,\"type\":\"subscribe_entities\",\"entity_ids\":[...]}\n"
     "                server:  one snapshot of every listed device   {\"event\":{\"a\":{...}}}\n"
     "                         then only what changes                {\"event\":{\"c\":{...}}}\n"
     "\n"
     "4.  Listen      {\"id\":2,\"type\":\"subscribe_events\",\"event_type\":\"homemesh_layout\"}\n"
     "                server:  layouts sent from the designer\n"
     "\n"
     "5.  Control     {\"id\":N,\"type\":\"call_service\",\"domain\":\"light\",\"service\":\"turn_on\",\n"
     "                 \"target\":{\"entity_id\":\"light.bed_light\"},\"service_data\":{\"brightness\":180}}")
para("Message 3 is how state comes in, message 5 is how commands go out, and message 4 is how the panel is "
     "reconfigured. There is nothing else and no second connection.")
figure(IMG / "paper_sequence.png", 9, "Message sequence between the panel and Home Assistant.")
para("<b>Availability</b> is the connection itself. The old MQTT version used a broker feature called a Last "
     "Will and Testament for this; the WebSocket API needs none. If the socket drops, the panel greys its tiles "
     "out immediately, and a 15-second heartbeat catches a connection that dies silently.")
h2("10.2 Task structure and interface responsiveness")
para("The graphics library must be serviced about every 5 ms to read touches, animate and redraw. Anything "
     "that blocks for longer causes visible stutter — and network operations routinely block for hundreds of "
     "milliseconds.")
para("So the network client runs as a separate FreeRTOS task on core 0 and the interface on core 1. The "
     "network task never touches the interface: it turns each change into a <c>PanelEvent</c> and puts it on a "
     "queue, and the interface drains the queue on each pass of its loop. A network operation that blocks for "
     "half a second has no effect on the screen; pages still swipe and sliders still track the finger.")
callout("Detail worth knowing: task priority",
        ["The network task runs at FreeRTOS priority 0, not 1. At 1 it starves the idle task on core 0, which "
         "is responsible for feeding the hardware watchdog; the watchdog concludes the system has hung and "
         "restarts it. Symptom: a panel that reboots every few seconds for no visible reason."])
h2("10.3 Memory strategy")
para("Memory shaped almost every implementation decision.")
bullets([
    "<b>Program storage.</b> The default 1.3 MB application partition was 92% full. A 3 MB partition brought "
    "it to 38%. The cost is the over-the-air update slot, which this project does not use.",
    "<b>Graphics memory</b> is a fixed 40 KB pool. Running out does not give an error message; it freezes or "
    "blanks the screen. The panel logs its headroom at every boot (about 15 KB free).",
    "<b>Static memory is the real ceiling.</b> The build report says about 37% of RAM is used, but the "
    "linker's data segment is the true limit, and the firmware sits about 1.5 KB below it: a 48 KB graphics "
    "pool, or even one extra 2 KB buffer, failed to link. Large buffers are shared, and growth must come from "
    "the heap (about 160 KB free).",
    "<b>Nothing large on the network task's stack.</b> A 4 KB JSON buffer declared inside a function "
    "overflowed the task's 12 KB stack, because the compiler inlined it into the handler that runs for every "
    "message. Big buffers now live in static memory.",
    "<b>Every message is classified on its raw bytes before parsing,</b> and each kind is parsed into a "
    "fixed-size buffer that keeps only the fields the panel uses. Memory use cannot grow with whatever Home "
    "Assistant decides to send.",
    "<b>Layout strings live in a fixed 2,560-byte arena,</b> so a layout can never take memory the interface "
    "needs. The current layout uses 275 bytes; a maximal one would use 2,320.",
])
h2("10.4 Touch coordinate handling")
para("The panel is physically portrait (240 × 320) but the interface is landscape (320 × 240). The touch "
     "controller still reports coordinates in the physical orientation, so they are rotated in software in "
     "the touch driver. Without this, touches land mirrored or on the wrong axis — which looks like broken "
     "hardware but is purely a rotation mismatch.")
h2("10.5 Receiving and storing a layout")
numbered([
    "The designer fires a <c>homemesh_layout</c> event in Home Assistant; Home Assistant relays it to every "
    "subscriber, including the panel. A 700-byte layout arrives as a 955-byte message.",
    "The panel recognises it from the raw message, before parsing, and parses it on its own path into a "
    "larger buffer, so the everyday path stays small.",
    "It validates the whole layout <b>before storing anything</b>: limits, rooms, kinds, string sizes. A "
    "rejected layout is reported with its reason on the serial log, and the panel keeps running the layout it "
    "has.",
    "An accepted layout is written to NVS and the board restarts into it. This is deliberate: rebuilding a "
    "whole screen in place fragments the 40 KB graphics pool, where running out freezes the screen; a restart "
    "takes a few seconds and always works.",
    "At every boot the panel reads NVS. If nothing is stored, or it does not validate, it falls back to the "
    "layout compiled into the firmware, so a board never pushed to still comes up working.",
])
para("Limits enforced by the firmware and the designer: <b>12 devices, 6 rooms, 6 tiles per page, 8 "
     "automations</b>. Pages do not scroll, so a seventh tile on a page would be unreachable.")
h2("10.6 The user interface")
para("A header shows the page title, that room's temperature and humidity, and a connection dot (green: "
     "online; amber: Wi-Fi but no server; red: no Wi-Fi). Pages swipe sideways: a Home page of favourites, then "
     "up to six more pages — by default one per room, but any page the designer names — each with up to six "
     "tiles in two columns. Swiping up opens an Automations screen for scenes, scripts and automations.")
figure(IMG / "ui_tile_detail.png", 15, "The Home page with lights on, and the detail panel for one light. "
       "Rendered by the panel's own interface code (Section 11.2), pixel for pixel.")
bullets([
    "<b>A tile</b> shows the device name, its icon, its state, and for dimmable lights a brightness bar. A "
    "chevron opens the detail panel: on/off, a brightness slider and a colour wheel.",
    "<b>On state:</b> the tile lifts a shade, the icon turns amber in a tinted disc, and the reading turns "
    "amber. A dimmed light shows its level, such as 72%, rather than \"ON\". Six solid amber tiles on a page "
    "was noise; the lift is quieter and still unmistakable.",
    "<b>Long names</b> are cut to one line with \"...\". The designer shows the exact cut while you type.",
    "<b>Icons</b> are 16 Material Design glyphs compiled into fonts, so they cost flash, not graphics memory.",
    "<b>Fills are flat.</b> Gradients were tried and removed: the 16-bit display has only 32 levels of red and "
    "blue, so a subtle gradient showed as visible, tinted bands.",
])
figure(FIG.parent / "paper" / "fig_ui_pages_lvgl.png", 13.5, "Every page of the current layout: Home and the "
       "three room pages. Rendered by the panel's own interface code.")
h2("10.7 Themes")
para("Nine colours define the look: background, tile, on, accent, text, muted text, and the three link "
     "colours. A layout may carry any of them; a missing colour keeps its default, and the designer only sends "
     "colours that differ. Six presets are built in, including a light theme (Daylight) and a high-contrast "
     "one.")
figure(IMG / "theme_presets.png", 12.5, "The six theme presets, each on the Home page, a room page with "
       "compact tiles, and the detail panel on and off. All rendered by the panel's own interface code.")
h2("10.8 A panel that explains its own failures")
para("Two failures used to be silent. If the graphics library ran out of memory it froze the screen; it now "
     "prints the exact file and line with the memory state, then restarts. Wi-Fi failures are now logged with "
     "the driver's reason code — for example <c>reason 201 (NO_AP_FOUND)</c> means the panel cannot see the "
     "network at all. That one line is what cracked the network fault in Section 13.4.")

# ======================================================================
#  11. DESIGNER
# ======================================================================
h1("11. The HomeMesh Designer (Desktop App)", new_page=True)
para("The designer is the PC half of the split architecture: it discovers devices, lets the user design the "
     "panel, and sends the result to it. It replaced a simpler Python tool on 20 October.")
figure(IMG / "designer_app.png", 16, "The designer connected to Home Assistant. Left: every entity, by room. "
       "Centre: the simulated panel and its log. Right: pages, names and colours.")
h2("11.1 What it does")
bullets([
    "<b>Connects to Home Assistant</b> and lists every entity the panel can drive, grouped by room. Lights, "
    "switches, fans and helper switches are controllable; scenes, scripts and automations are actions; sensors "
    "are kept only if they measure temperature or humidity.",
    "<b>Pick and arrange:</b> tick a device to place it, star it for the Home page, choose its icon, drag tiles "
    "to reorder, and move whole pages up and down.",
    "<b>Pages are yours, not just Home Assistant's rooms.</b> Each device has a page chooser that starts at its "
    "room; put it on any page instead, make new pages, rename them (\"Living Room\" can be \"Lounge\" on the "
    "panel), or drag a tile onto another page. A device with no room in Home Assistant is placed by choosing a "
    "page for it. The temperature and humidity shown in a page's header can be chosen the same way. To the "
    "panel a page is just a name, so none of this needed a firmware change.",
    "<b>Rename</b> any device for the panel. Each name shows exactly how the panel will cut it "
    "(\"Ceiling...\") and warns about characters the panel's font cannot draw, such as accents.",
    "<b>Colours:</b> six presets or any of the nine colours by hand.",
    "<b>Live states:</b> the simulated panel shows real device states, and tapping it really switches the "
    "devices, using the same commands the firmware sends.",
    "<b>Save to project</b> writes the layout file and the compiled fallback; <b>Push to panel</b> sends it "
    "through Home Assistant. Unsaved changes are flagged and can be reverted.",
    "Light, dark or automatic appearance for the app itself. Ships as a single portable .exe.",
])
h2("11.2 The simulator is the real interface")
para("The preview is not a drawing of the panel. It is the panel's own interface code — "
     "<c>dashboard.cpp</c>, LVGL 8.3, the icon fonts and the layout validator — compiled to WebAssembly with "
     "Emscripten against the firmware's own configuration. Same pixels, same fonts, same 16-bit colour, the "
     "same 40 KB graphics memory, and the same validator: if the simulator rejects a layout, the panel would "
     "too, with the same message. Nothing in the firmware source is copied or changed for it.")
para("Applying a layout restarts the simulated panel, just as it restarts the real one. The name hints come "
     "from LVGL itself: an earlier approach that predicted the cut in JavaScript got three of eight test names "
     "wrong, because LVGL wraps at word boundaries, rounds each character to whole pixels and applies kerning.")
h2("11.3 How the designer is tested")
bullets([
    "The layout format is written once, in Python (<c>ha_discover.py</c>). The app's version must produce "
    "<b>byte-identical</b> output from a stored Home Assistant snapshot, checked automatically.",
    "Name cuts are checked against cuts the firmware actually rendered.",
    "A live test drives a real light through Home Assistant and back.",
    "Every figure of the panel in this document was rendered by the same simulator.",
])

# ======================================================================
#  12. SYSTEM OPERATION
# ======================================================================
h1("12. System Operation", new_page=True)
h2("12.1 Network configuration")
para("The demonstration runs on a Wi-Fi hotspot from the laptop, not on the college network, which breaks the "
     "system three ways: client isolation stops the panel reaching the laptop, a captive login page cannot be "
     "completed by a microcontroller, and local name lookup is filtered. The hotspot makes the demo independent "
     "of the venue.")
table(["Property", "Value", "Why it matters"], [
    ["Server address", "192.168.137.1", "Windows always gives the hotspot host this address, so it can be "
     "written into the firmware."],
    ["Frequency band", "2.4 GHz only", "The ESP32 has no 5 GHz radio. On this laptop the Wi-Fi card itself "
     "must be pinned to 2.4 GHz (Section 13.4)."],
    ["Security", "WPA2", "WPA3 is unreliable on this chip with the Arduino framework."],
    ["Home Assistant", "Docker on the laptop", "Port 8123 on the hotspot address; Docker's own firewall rule "
     "covers it."],
], [3.2, 3.6, 9.2], size=9)
para("The hotspot is used for development too, not only for the demo. Had development used another network, "
     "the server address would change on the day, and because it is compiled into the firmware that would mean "
     "reflashing on the morning of the review.")
h2("12.2 Devices on the panel")
rows = []
for i, a in enumerate(areas):
    dev = [ent_desc(e) for e in ents if e["area"] == i]
    sen = [("Humidity" if s["hum"] else "Temperature") for s in sensors if s["area"] == i]
    rows.append([a, "\n".join(dev) or "—", ", ".join(sen) or "none"])
favs = [e["name"] for e in ents if e.get("fav")]
para(f"The layout in the repository (<c>include/ha_layout.json</c>, read when this document was built) has "
     f"{len(ents)} devices in {len(areas)} rooms, with {len(favs)} favourites on the Home page: "
     f"{', '.join(favs)}. The panel itself shows whatever was last pushed to it, and it prints which layout "
     f"it loaded at boot.")
table(["Room page", "Devices", "Sensors"], rows, [3.2, 9.0, 3.8], size=9)
para("Rubric minimum, all met: at least three lights with at least one brightness slider (every light has "
     "one), a switch, a temperature and humidity readout, and a connection indicator.")
h2("12.3 Changing what the panel shows")
numbered([
    "Open the designer. It connects to Home Assistant and lists the devices.",
    "Tick, star, rename, reorder and colour; choose or create a page for any device. The simulated panel "
    "updates as you go.",
    "Press <b>Push to panel</b>. The panel restarts and comes back with the new layout in a few seconds.",
    "Optionally <b>Save to project</b>, so the compiled fallback matches.",
])

# ======================================================================
#  13. TESTING AND RESULTS
# ======================================================================
h1("13. Testing and Results", new_page=True)
para("The system was tested end to end against a real Home Assistant installation (version 2026.9, in Docker "
     "on a Windows 11 laptop) whose demonstration integration provides lights, switches, fans and sensors — 123 "
     "entities in all. Server-side figures were measured with a script on the laptop; panel figures come from "
     "the panel's serial log and the build report. Measured 18 October.")
h2("13.1 Message size and latency")
table(["Measurement", "Result"], [
    ["Full state dump, get_states (123 entities)", "57,671 bytes"],
    ["Average size per entity", "469 bytes"],
    ["Panel's initial state, subscribe_entities (7 entities)", "2,403 bytes, one message"],
    ["Change message: brightness only / light off", "144 bytes / 262 bytes"],
    ["Layout delivery (700-byte layout)", "955-byte message"],
    ["Snapshot delivered after subscribing (server)", "11.5 ms"],
    ["Command to state-change message, median of 10 (server)", "3.2 ms"],
    ["Panel round-trip time over the hotspot, mean of 6", "40 ms, 0% loss"],
], [10.0, 6.0], size=9.5)
figure(IMG / "paper_message_sizes.png", 11, "Message sizes: what the panel avoids, and what it handles.")
para("The initial state is 24 times smaller than the full dump and, more importantly, depends on the layout "
     "rather than the size of the house. A command produces its state change in milliseconds — far below the "
     "point where a touch response feels delayed.")
h2("13.2 Resource usage")
table(["Resource", "Used / available"], [
    ["Flash (application partition)", "1,198,037 / 3,145,728 bytes (38.1%)"],
    ["Static RAM", "123,024 bytes; about 1.5 KB below the linker limit"],
    ["Graphics pool, free after the first full draw", "15,652 / 40,960 bytes"],
    ["Free heap after start-up (lowest)", "161,028 bytes (158,632)"],
    ["Network task stack, free after login", "8,272 / 12,288 bytes"],
    ["Icon fonts (16 glyphs, two sizes)", "about 7.6 KB of flash, 0 bytes of graphics memory"],
], [8.0, 8.0], size=9.5)
h2("13.3 Functional verification")
table(["Test", "Result"], [
    ["Connect and log in to Home Assistant over the hotspot", "Pass"],
    ["Receive the initial state of every device in one message", "Pass"],
    ["Switch a device from the panel; it changes in Home Assistant", "Pass"],
    ["Change a device in Home Assistant; the panel follows", "Pass"],
    ["Set brightness and colour from the detail panel", "Pass"],
    ["Show temperature and humidity on the right room page", "Pass"],
    ["Swipe pages during network traffic", "Pass — no stutter"],
    ["Stop the server; the panel shows it", "Pass — tiles greyed out, no stale values"],
    ["Push a layout from the designer; the panel stores it and restarts into it", "Pass — on the board, "
     "17 October"],
    ["Reject an invalid layout and keep the old one", "Pass — reason printed on the log"],
    ["Boot with nothing stored; fall back to the compiled layout", "Pass — on the board, 16 October"],
    ["Designer: tap the simulated panel in Live mode; the real light switches", "Pass — 20 October"],
    ["Designer layout output identical to the reference tool", "Pass — automated test"],
], [11.0, 5.0], size=9)
h2("13.4 Field case study: the connection that kept dropping")
para("For days the panel would randomly show no connection, tiles read \"--\", and faded tiles seemed to "
     "linger between pages. It looked like flaky firmware. It was two problems stacked on each other.")
numbered([
    "<b>A design flaw that amplified faults.</b> The panel fetched each device's starting state with an HTTP "
    "request made from inside the WebSocket's own message handler. When those requests were slow — up to 21 "
    "seconds for seven devices — the connection's heartbeat could not run, so the library dropped the "
    "connection, reconnected and did it all again. <b>Fix:</b> take the starting state from the subscription "
    "on the existing connection. Nothing can block the handler any more, and the HTTP library and 34 KB of "
    "flash went with it.",
    "<b>The root cause was the laptop.</b> The panel then failed to join the network at all, and the new logging "
    "said why: <c>reason 201 (NO_AP_FOUND)</c>. The laptop's Wi-Fi card has a single radio, so the hotspot "
    "must share the channel of the laptop's own Wi-Fi connection. Whenever the router moved the laptop to "
    "5 GHz, the hotspot silently moved with it — whatever its own setting said — and became invisible to the "
    "2.4 GHz-only panel. The router changes band as it pleases, which is why it looked random. <b>Fix:</b> "
    "pin the card to 2.4 GHz in Device Manager. The panel joined at −35 dBm with no packet loss, and the "
    "preflight script now checks the setting.",
])
para("The lesson worth presenting: a flaky network alone would have cost a few readings; blocking inside the "
     "message handler turned each slow reply into a dropped connection. And logging that names its causes "
     "turned a days-long mystery into a one-session diagnosis.")
h2("13.5 Other challenges and their resolution")
table(["Challenge", "Resolution"], [
    ["Board carried different display and touch controllers than documented",
     "Identified both by testing; confined hardware settings to two driver files."],
    ["Panel rebooted every few seconds under network load",
     "Network task starved the idle task and tripped the watchdog. Priority lowered to 0."],
    ["Touches landed in the wrong place", "Touch reports physical orientation; rotation added in the driver."],
    ["Program too big for the application partition", "Switched to a 3 MB partition, dropping the unused "
     "over-the-air slot."],
    ["Crash at start-up inside an unrelated print statement",
     "A 4 KB buffer inlined onto the network task's stack. Moved to static memory."],
    ["Layout pushes vanished with no error",
     "A fixed-size JSON filter silently dropped a key that did not fit. Pushes are now recognised from the raw "
     "message, before parsing."],
    ["Larger graphics pool would not link", "Static memory, not total RAM, is the limit. Buffers shared; growth "
     "must come from the heap."],
    ["Gradients showed visible bands", "16-bit colour has 32 levels of red and blue. Fills made flat."],
    ["Long names wrapped into the icon", "Cut to one line with \"...\"; the designer shows the cut."],
    ["Light theme: names vanished on lit tiles, invisible close button",
     "LVGL's default theme forces white text on checked buttons; colours now set explicitly. Found by "
     "rendering every preset in the simulator."],
    ["Panel joined Wi-Fi but could not reach the server",
     "The firmware had the house network's name, not the hotspot's. Corrected."],
    ["Most Home Assistant devices invisible", "The reason for leaving MQTT (Section 5.1)."],
], [6.4, 9.6], size=9)

# ======================================================================
#  14. FUTURE SCOPE
# ======================================================================
h1("14. Future Scope")
bullets([
    "<b>Encrypted connection and token provisioning.</b> The connection is unencrypted and the access token is "
    "compiled in — acceptable on an isolated hotspot, not in a real home.",
    "<b>Multi-cell widgets:</b> a two-wide tile with its controls inline instead of behind the chevron. Needs "
    "real placement in the layout format, the firmware and the designer.",
    "<b>Acknowledged layout delivery,</b> so the designer can confirm the panel took a layout. Today only the "
    "panel's serial log confirms it.",
    "<b>Use the microSD card</b> for logging or images, once it has been tested.",
    "<b>A dedicated ESP32 access point</b> so the panel does not depend on a laptop hotspot. Written, not yet run.",
    "<b>Remove the legacy MQTT code</b> to shrink what can break.",
    "More device types — media players, covers, thermostats — multiple panels with their own layouts, and a "
    "screen sleep mode.",
    "End-to-end latency measured from touch to a physical device.",
])

# ======================================================================
#  15. CONCLUSION
# ======================================================================
h1("15. Conclusion")
para("The project set out to build a wall-mounted touchscreen that controls a home automation system entirely "
     "over the local network, on hardware costing a small fraction of a commercial panel. That has been met and "
     "demonstrated against a real installation: the panel switches, dims and colours lights, shows temperature "
     "and humidity, follows changes made anywhere else in the house, and stays responsive while the network is "
     "busy, with no cloud service involved.")
para("Three decisions account for most of its value. <b>Replacing MQTT with the Home Assistant WebSocket "
     "API</b> changed device coverage from a subset that depended on how each device was set up to every "
     "device in the system. <b>Recognising that on-device discovery was impossible</b> and splitting the work — "
     "the PC discovers, the panel subscribes — keeps every message small: 2,403 bytes where the naive design "
     "needed 57,671 for a small house. And <b>making the layout data</b> means the panel is redesigned in a "
     "desktop app and updated over the network in seconds, with a preview that is the panel's own code.")
para("The field case study adds a fourth lesson: firmware that reports why it fails is worth more than "
     "firmware that tries to hide it.")

# ======================================================================
#  16. REFERENCES
# ======================================================================
h1("16. References")
h2("16.1 Published literature")
para("Each entry was checked against the Crossref DOI registry. The same list, with in-text citations, is in "
     "the IEEE paper.", size=9.5, color=GREY)
src = (HERE / "paper" / "build_paper.py").read_text(encoding="utf-8")
m = re.search(r"^REFS = \{.*?^\}", src, re.S | re.M)
ns = {}
exec(m.group(0), ns)
REFS = ns["REFS"]
# order: the order this document's literature table cites them in
order = ["reis", "swamy", "russo", "madadi", "stolojescu", "verma", "tse", "veigl", "gupta", "donta",
         "bayilmis", "lalhriatpuii", "mishra", "kaur", "ajayi", "yamin", "demir"]
assert sorted(order) == sorted(REFS), "reference list out of step with the paper"
for i, k in enumerate(order, 1):
    p = para(f"[{i}]  " + REFS[k], size=9, after=3)
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
h2("16.2 Technical documentation")
tech = [
    "Home Assistant. WebSocket API. developers.home-assistant.io/docs/api/websocket",
    "Home Assistant. Architecture: entities and the entity registry. developers.home-assistant.io/docs/architecture_index",
    "Espressif Systems. ESP32-WROOM-32 Datasheet.",
    "Espressif Systems. ESP-IDF FreeRTOS reference: task priorities and the task watchdog.",
    "LVGL. Light and Versatile Graphics Library, version 8.3 documentation. docs.lvgl.io/8.3",
    "Bodmer. TFT_eSPI: TFT graphics library for ESP8266 and ESP32. github.com/Bodmer/TFT_eSPI",
    "Links2004. arduinoWebSockets: WebSocket server and client for Arduino. github.com/Links2004/arduinoWebSockets",
    "B. Blanchon. ArduinoJson: JSON library for embedded C++. arduinojson.org",
    "Sitronix Technology. ST7789V TFT controller/driver datasheet.",
    "Goodix Technology. GT911 capacitive touch controller datasheet.",
    "OASIS. MQTT version 3.1.1 specification.",
    "PlatformIO. Espressif 32 platform documentation. docs.platformio.org",
    "Emscripten. Compiling C and C++ to WebAssembly. emscripten.org/docs",
    "Electron. Electron documentation. electronjs.org/docs",
]
for i, t in enumerate(tech, len(order) + 1):
    p = para(f"[{i}]  {t}", size=9, after=3)
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)

# ======================================================================
#  APPENDIX A — PRESENTATION PLAN
# ======================================================================
h1("Appendix A — Suggested Presentation Plan", new_page=True)
para("Twenty slides. The content for each is drawn from the section shown, so nothing needs writing from "
     "scratch. Figure files are in <c>docs/figures/</c> and <c>docs/paper/</c>.")
table(["#", "Slide", "What to put on it", "Source / figure"], [
    ["1", "Title", "Project title, course code, team members with register numbers, guide, date.", "Title page"],
    ["2", "Introduction", "Phone apps are slower than the switch they replace; a panel is always on and always "
     "in one place. No cloud.", "§2"],
    ["3", "Literature survey", "The table, shortened to one line per row.", "§2.2"],
    ["4", "Problem statement", "The six conditions, one short line each.", "§3"],
    ["5", "Objectives", "The eight objectives, one line each.", "§4"],
    ["6", "Existing systems", "The three approaches and their limitations.", "§5"],
    ["7", "Why not MQTT", "<b>Key slide.</b> Same Zigbee bulb: reachable or invisible depending only on how it "
     "was added. Then the three designs table.", "§5.1, §5.2"],
    ["8", "Proposed system", "WebSocket API: every device, one command.", "§6"],
    ["9", "The memory wall", "<b>Key slide.</b> 57,671 bytes vs a 40 KB pool. PC discovers, panel subscribes. "
     "2,403 bytes.", "§6.1, paper_message_sizes.png"],
    ["10", "Overall architecture", "Full width. Setup path first, then runtime.", "arch_overall.png"],
    ["11", "Hardware", "Diagram plus the components table. Mention the controllers had to be identified by "
     "testing.", "arch_hardware.png, §8"],
    ["12", "Software", "The two cores and the queue between them.", "arch_software.png"],
    ["13", "Protocol", "The five messages. State in, commands out, layouts in.", "§10.1, paper_sequence.png"],
    ["14", "Layout as data", "Designer → Home Assistant event → validate → NVS → restart. No reflash.",
     "§6.2, §10.5"],
    ["15", "The interface", "Tiles, on state, detail panel; the themes.", "ui_tile_detail.png, "
     "theme_presets.png"],
    ["16", "The designer", "Screenshot. The preview is the panel's own code in WebAssembly; Live mode switches "
     "real devices.", "designer_app.png, §11"],
    ["17", "Results", "Message size table and resource table.", "§13.1, §13.2"],
    ["18", "Field case study", "<b>Best story.</b> Random disconnects → blocking handler + single-radio Wi-Fi "
     "on 5 GHz. Found because the panel logs why it fails.", "§13.4"],
    ["19", "Future scope", "Encryption, multi-cell tiles, acknowledged push, microSD, dedicated access point.",
     "§14"],
    ["20", "Conclusion + demo", "The three decisions. Then the live demonstration.", "§15"],
], [0.8, 3.0, 8.4, 3.8], size=8.5)
h2("Figure files")
table(["Figure", "File"], [
    ["Overall architecture", "docs/figures/arch_overall.png"],
    ["Hardware architecture", "docs/figures/arch_hardware.png"],
    ["Software architecture", "docs/figures/arch_software.png"],
    ["Block diagram", "docs/figures/img/paper_block_diagram.png"],
    ["Message sequence", "docs/figures/img/paper_sequence.png"],
    ["Message sizes chart", "docs/figures/img/paper_message_sizes.png"],
    ["Panel: Home page and detail panel", "docs/figures/img/ui_tile_detail.png"],
    ["Panel: every page", "docs/paper/fig_ui_pages_lvgl.png"],
    ["Panel: six themes", "docs/figures/img/theme_presets.png"],
    ["Designer app", "docs/figures/img/designer_app.png"],
], [6.0, 10.0], size=9)
h2("Presentation notes")
bullets([
    "Slides 7, 9 and 18 carry the engineering argument. If time is short, cut elsewhere.",
    "Use the architecture figures full width; shrunk, their labels become unreadable.",
    "The strongest live moment: change a light in Home Assistant's web page and let the audience watch the "
    "panel follow — two-way sync in a way no slide can show. The second strongest: rename a device in the "
    "designer, press Push, and watch the panel restart with the new name.",
    "Panel images in this document are rendered by the panel's own code, so they are exact — but photograph "
    "the real panel for at least one slide, so nobody thinks it is a mock-up.",
])
h2("Anticipated questions")
table(["Question", "Answer"], [
    ["Why not a Raspberry Pi with a browser?", "Cost, boot time and power. This starts in a couple of seconds "
     "and draws a fraction of the power, with no operating system to maintain on a wall."],
    ["What happens if Home Assistant goes down?", "The connection drops, the panel sees it immediately and greys "
     "the tiles out. It never shows stale values that look live."],
    ["Are these real devices?", "They come from Home Assistant's demonstration integration. The protocol is "
     "identical to a physical bulb — the panel cannot tell the difference, and real hardware needs no firmware "
     "change."],
    ["How do you change what is on the panel?", "In the designer: tick, star, rename, reorder, pick colours, "
     "Push. The panel restarts with the new layout in seconds. No cable, no reflash."],
    ["Why does it restart when a layout arrives?", "Rebuilding a screen in place fragments the 40 KB graphics "
     "memory, and running out freezes the screen. A restart takes seconds and always works."],
    ["Is the designer's preview a mock-up?", "No — it is the panel's own interface code compiled to "
     "WebAssembly. Same fonts, same pixels, same memory limit, same validator."],
    ["What was the hardest bug?", "The random disconnects in §13.4: a blocking request in the message handler, "
     "plus the laptop's single-radio Wi-Fi card dragging the hotspot onto 5 GHz."],
    ["Why LVGL 8.3 and not 9?", "Version 9 changed the driver interfaces, and nearly every example for this "
     "board is written against 8.3."],
    ["How did you know which display controller it has?", "By testing. This board ships with different "
     "controllers in different batches."],
    ["How many devices can it handle?", "12 devices, 6 rooms, 6 per page, 8 automations. Pages do not scroll, "
     "so a seventh tile would be unreachable."],
    ["Is it secure?", "On an isolated hotspot, yes enough. The connection is unencrypted and the token is in "
     "the firmware, so a real home would need TLS — listed as future work."],
    ["What is the microSD card for?", "Nothing yet; it has not been tested. Logging is the likely use."],
], [5.0, 11.0], size=9)

# ======================================================================
#  APPENDIX B — INTERNAL NOTE
# ======================================================================
h1("Appendix B — Internal Note (Not for Presentation)", new_page=True)
callout("This appendix is for the team only",
        ["Do not show this page, read it aloud, or put it on a slide. It records what is not finished, so that "
         "nothing is claimed during the review that cannot be demonstrated."])
h2("What is genuinely working")
para("All of these can be demonstrated live and claimed without qualification:")
bullets([
    "The panel controls a real Home Assistant installation over the local network, with no cloud.",
    "Two-way sync: a change in Home Assistant appears on the panel, and the panel's taps change Home Assistant.",
    "Brightness and colour control on every light; temperature and humidity; the connection indicator.",
    "Smooth page swiping during network traffic.",
    "Layouts pushed over the network and stored on the panel, with no reflash — proven on the board on "
    "17 October.",
    "The designer: discovery, picking, favourites, icons, reordering, renaming, themes, the simulated panel, and "
    "Live mode switching real devices.",
])
h2("What must not be claimed")
table(["Do not say", "Say instead"], [
    ["That MQTT has been removed", "The MQTT code is still in the project, unused, kept as a fallback."],
    ["That the microSD card is used or works", "It is fitted but has not been tested."],
    ["That the connection is encrypted", "It is not; that is future work (§14)."],
    ["That the designer knows the panel received a layout", "There is no acknowledgement. \"Pushed\" means "
     "Home Assistant accepted it; the panel's serial log is the only confirmation."],
    ["That latency was measured from touch to bulb", "Latency was measured at the server (3.2 ms) and as "
     "network round trip (40 ms)."],
    ["That the ESP32 access point works", "It is written but has never been run."],
    ["That tiles can be resized or placed freely", "Tiles are placed by order in a fixed two-column grid."],
], [6.0, 10.0], size=9)
h2("Check before the review: firmware on the board")
callout("The board is running the 20 October firmware", [
    "Flashed on 20 October with the fixes for long names and the light (Daylight) theme; the boot log matched "
    "the healthy one. Check the names and the Daylight theme on the real screen once before the review.",
    "The new designer's Push sends the same event as the old tool (tested byte for byte), but nobody has yet "
    "watched a push from <i>the new designer</i> arrive on the board. Do it once with the serial monitor open "
    "and look for <c>[layout] push received</c> then <c>[layout] loaded from NVS</c>.",
    "The Automations screen needs at least one scene, script or automation in Home Assistant. The demo "
    "installation has none, so do not promise it unless one is added.",
])
h2("Two quirks that may be noticed, and how to answer")
table(["If someone notices", "The answer"], [
    ["Every light is listed as colour-capable (RGB), not dimmable", "Correct, not a fault. RGB means on/off "
     "plus brightness plus colour, so every light gets a slider. The requirement was at least one."],
    ["Bedroom and Kitchen show no temperature or humidity", "The demo installation has one temperature and "
     "humidity pair, in the Living Room. Showing nothing beats showing an invented reading."],
], [6.0, 10.0], size=9)
h2("Before the presentation")
code("powershell -ExecutionPolicy Bypass -File tools\\preflight.ps1")
bullets([
    "Run it from cold. It checks the Wi-Fi card is pinned to 2.4 GHz, the hotspot is up, the firewall path, "
    "that Home Assistant is running and reachable on 192.168.137.1, the token, and every device in the layout. "
    "It names the fix for anything that fails.",
    "If the panel does not connect, read its serial log first. <c>reason 201 (NO_AP_FOUND)</c> means it cannot "
    "see the network — almost certainly the hotspot is on 5 GHz: check Band Selection and toggle the hotspot.",
    "If the venue Wi-Fi is unusable, the hotspot still needs a network to share: join the laptop to a phone "
    "hotspot set to 2.4 GHz first.",
    "If something is badly wrong: <c>git reset --hard v0.4-firmware-freeze</c>, or flash "
    "<c>../homemesh_ha_firmware_18oct_full.bin</c> (restores the whole board; push the layout again "
    "afterwards), or restore from GitHub.",
])

for s in doc.sections:
    page_number_footer(s)

doc.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)}  ({FIGNO[0]} figures)")

# ----------------------------------------------------------------------
#  Word: fill in the contents page, save, export a PDF beside it
# ----------------------------------------------------------------------
if "--no-word" not in sys.argv and sys.platform == "win32":
    pdf = OUT.with_suffix(".pdf")
    ps = f"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false; $word.DisplayAlerts = 0
try {{
  $d = $word.Documents.Open('{OUT}', $false, $false)
  $d.Repaginate()
  foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
  $d.Repaginate()
  foreach ($t in $d.TablesOfContents) {{ $t.UpdatePageNumbers() }}
  Write-Host ('pages: ' + $d.ComputeStatistics(2) + '   words: ' + $d.ComputeStatistics(0))
  $d.Save()
  $d.ExportAsFixedFormat('{pdf}', 17)
  $d.Close($false)
}} finally {{ $word.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($word) }}
"""
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
                       capture_output=True, text=True)
    print((r.stdout + r.stderr).strip() or "Word step finished")
